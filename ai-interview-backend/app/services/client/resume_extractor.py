"""统一简历文本提取器 — PDF / DOCX / PPTX → 纯文本。

非 AI 业务工具，供 ResumeService 上传解析使用。文本提取产出与格式无关的
纯文本，后续统一交给 ai_service.parse_resume() 做 LLM 解析。
"""
from pathlib import Path
from app.exceptions.http_exceptions import ValidationError

# 支持的简历格式（现代 OOXML，纯 Python 库解析）
SUPPORTED_EXTS = {".pdf", ".docx", ".pptx"}

# 老式格式，提示另存为
LEGACY_EXTS = {".doc", ".ppt"}


def validate_resume_ext(file_name: str) -> None:
    """校验文件名后缀（两档提示）。

    - .pdf / .docx / .pptx：通过
    - .doc / .ppt：提示另存为现代格式
    - 其他：仅支持提示
    """
    ext = Path(file_name).suffix.lower()
    if ext in SUPPORTED_EXTS:
        return
    if ext in LEGACY_EXTS:
        raise ValidationError(message="不支持 .doc/.ppt 格式，请将文档另存为 .docx 或 .pptx 后上传")
    raise ValidationError(message="仅支持 PDF / Word (.docx) / PPT (.pptx) 格式文件")


def extract_resume_text(file_path: str, file_name: str) -> str:
    """按扩展名分派提取纯文本；先做 magic-byte 校验文件真实格式。"""
    validate_resume_ext(file_name)
    ext = Path(file_name).suffix.lower()
    _check_magic(file_path, ext)

    if ext == ".pdf":
        text = _extract_pdf_text(file_path)
    elif ext == ".docx":
        text = _extract_docx_text(file_path)
    else:
        text = _extract_pptx_text(file_path)

    if not text.strip():
        raise ValidationError(message="无法从文件中提取文本内容")
    return text


def _check_magic(file_path: str, ext: str) -> None:
    """校验文件头魔数，拒绝改名冒充的文件。"""
    with open(file_path, "rb") as f:
        head = f.read(5)
    if ext == ".pdf":
        if not head.startswith(b"%PDF-"):
            raise ValidationError(message="文件内容与格式不匹配")
    elif ext in {".docx", ".pptx"}:
        if not head.startswith(b"PK\x03\x04"):
            raise ValidationError(message="文件内容与格式不匹配")


def _extract_pdf_text(file_path: str) -> str:
    """从 PDF 提取文本（pdfplumber）。"""
    import pdfplumber
    text = ""
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text


def _extract_docx_text(file_path: str) -> str:
    """全面提取 docx：段落 + 表格 + 页眉页脚 + 文本框（w:txbxContent）。"""
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(file_path)
    parts: list[str] = []

    # 正文段落
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())

    # 表格（每行单元格用 | 分隔）
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            cells = [c for c in cells if c]
            if cells:
                parts.append(" | ".join(cells))

    # 页眉页脚
    for section in doc.sections:
        for part in (section.header, section.footer):
            for para in part.paragraphs:
                if para.text.strip():
                    parts.append(para.text.strip())

    # 文本框（走底层 OOXML w:txbxContent）
    for txbx in doc.element.body.iter(qn("w:txbxContent")):
        for t_para in txbx.iter(qn("w:p")):
            text = "".join(node.text or "" for node in t_para.iter(qn("w:t")))
            if text.strip():
                parts.append(text.strip())

    return "\n".join(parts)


def _extract_pptx_text(file_path: str) -> str:
    """逐页提取 pptx：文本框架 + 表格，组形状递归下钻；跳过备注。"""
    from pptx import Presentation

    prs = Presentation(file_path)
    pages: list[str] = []
    for slide in prs.slides:
        lines: list[str] = []
        _collect_slide_text(slide.shapes, lines)
        if lines:
            pages.append("\n".join(lines))
    return "\n\n".join(pages)


def _collect_slide_text(shapes, lines: list[str]) -> None:
    """递归收集一页内所有文本（含组形状嵌套）。"""
    for shape in shapes:
        if getattr(shape, "shapes", None):  # 组形状
            _collect_slide_text(shape.shapes, lines)
            continue
        if getattr(shape, "has_text_frame", False):
            for para in shape.text_frame.paragraphs:
                # para.text 包含 a:br（软换行）与 a:fld（域），a:br 在 python-pptx
                # 中表示为 \v（垂直制表符），转成 \n 以免电话号码/邮箱等被吞并
                text = para.text.replace("\v", "\n").strip()
                if text:
                    lines.append(text)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells]
                cells = [c for c in cells if c]
                if cells:
                    lines.append(" | ".join(cells))
