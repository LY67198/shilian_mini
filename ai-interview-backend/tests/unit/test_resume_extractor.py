"""resume_extractor — 简历多格式文本提取器单测"""
import pytest
from app.exceptions.http_exceptions import ValidationError
from app.services.client.resume_extractor import (
    extract_resume_text,
    validate_resume_ext,
)


@pytest.mark.unit
class TestValidateExt:
    def test_supported_passes(self):
        for name in ("r.pdf", "r.docx", "r.PPTX"):
            validate_resume_ext(name)  # 不应抛错

    def test_legacy_doc_ppt_friendly(self):
        for name in ("r.doc", "r.ppt"):
            with pytest.raises(ValidationError) as e:
                validate_resume_ext(name)
            assert "另存为" in e.value.detail

    def test_unknown_ext_generic(self):
        with pytest.raises(ValidationError) as e:
            validate_resume_ext("r.txt")
        assert "仅支持" in e.value.detail


@pytest.mark.unit
class TestMagicByte:
    def test_pdf_ext_zip_content_rejected(self, tmp_path):
        p = tmp_path / "fake.pdf"
        p.write_bytes(b"PK\x03\x04 fakezip")
        with pytest.raises(ValidationError) as e:
            extract_resume_text(str(p), "fake.pdf")
        assert "文件内容与格式不匹配" in e.value.detail

    def test_docx_ext_nonzip_rejected(self, tmp_path):
        p = tmp_path / "fake.docx"
        p.write_bytes(b"not a zip")
        with pytest.raises(ValidationError) as e:
            extract_resume_text(str(p), "fake.docx")
        assert "文件内容与格式不匹配" in e.value.detail


@pytest.mark.unit
class TestDocxExtraction:
    def _build_docx(self, path, include_header=True):
        from docx import Document
        doc = Document()
        doc.add_paragraph("Hello World")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A1"
        table.cell(0, 1).text = "B1"
        table.cell(1, 0).text = "A2"
        table.cell(1, 1).text = "B2"
        if include_header:
            doc.sections[0].header.paragraphs[0].text = "HeaderText"
        doc.save(str(path))

    def test_paragraph_table_header(self, tmp_path):
        p = tmp_path / "r.docx"
        self._build_docx(p)
        text = extract_resume_text(str(p), "r.docx")
        assert "Hello World" in text
        assert "A1 | B1" in text
        assert "A2 | B2" in text
        assert "HeaderText" in text
        assert text.index("Hello World") < text.index("A1 | B1")

    def _inject_textbox(self, path):
        """在 word/document.xml 末尾注入一个文本框（VML + w:txbxContent）。"""
        import zipfile
        with zipfile.ZipFile(str(path)) as z:
            names = z.namelist()
            data = {n: z.read(n) for n in names}
        xml = data["word/document.xml"].decode("utf-8")
        textbox_xml = (
            '<w:p><w:r><w:pict>'
            '<v:shape xmlns:v="urn:schemas-microsoft-com:vml">'
            '<w:txbxContent><w:p><w:r><w:t>TextBoxLine</w:t></w:r></w:p></w:txbxContent>'
            '</v:shape></w:pict></w:r></w:p>'
        )
        xml = xml.replace("</w:body>", textbox_xml + "</w:body>")
        data["word/document.xml"] = xml.encode("utf-8")
        with zipfile.ZipFile(str(path), "w", zipfile.ZIP_DEFLATED) as z:
            for n, content in data.items():
                z.writestr(n, content)

    def test_textbox_extracted(self, tmp_path):
        p = tmp_path / "tb.docx"
        self._build_docx(p, include_header=False)
        self._inject_textbox(p)
        text = extract_resume_text(str(p), "tb.docx")
        assert "TextBoxLine" in text
        assert text.count("TextBoxLine") == 1
        # 文本框内容排在最后（页眉页脚之后）
        assert text.index("TextBoxLine") > text.index("A2 | B2")

    def test_empty_docx_raises(self, tmp_path):
        from docx import Document
        p = tmp_path / "empty.docx"
        Document().save(str(p))
        with pytest.raises(ValidationError) as e:
            extract_resume_text(str(p), "empty.docx")
        assert "无法从文件中提取文本内容" in e.value.detail


@pytest.mark.unit
class TestPptxExtraction:
    def _build_pptx(self, path):
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        # 页 1：标题文本
        s1 = prs.slides.add_slide(prs.slide_layouts[1])
        s1.shapes.title.text = "Page1Title"
        # 页 2：表格 + 组形状嵌套文本
        s2 = prs.slides.add_slide(prs.slide_layouts[1])
        tbl = s2.shapes.add_table(
            rows=1, cols=2, left=Inches(1), top=Inches(1),
            width=Inches(4), height=Inches(1),
        )
        tbl.table.cell(0, 0).text = "A1"
        tbl.table.cell(0, 1).text = "B1"
        group = s2.shapes.add_group_shape()
        tb = group.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
        tb.text_frame.text = "NestedText"
        prs.save(str(path))

    def test_multi_slide_table_group(self, tmp_path):
        p = tmp_path / "r.pptx"
        self._build_pptx(p)
        text = extract_resume_text(str(p), "r.pptx")
        assert "Page1Title" in text
        assert "A1 | B1" in text
        assert "NestedText" in text
        assert text.index("Page1Title") < text.index("A1 | B1")
